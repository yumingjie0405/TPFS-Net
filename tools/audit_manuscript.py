"""Audit a TPFS-Net DOCX for result provenance and numbering consistency."""

import argparse
import json
import re
from pathlib import Path

from docx import Document


UNSUPPORTED_PATTERNS = {
    'unsupported_nw_ucla_result': re.compile(r'(NW[-–]?UCLA|Northwestern-UCLA).{0,80}97\.0', re.I),
    'unsupported_flops': re.compile(r'1\.65\s*(G?FLOPs?|G)', re.I),
    'unsupported_five_seed_robustness': re.compile(r'(五次|5\s*(次|个)随机|five\s+(runs|seeds))', re.I),
    'unsupported_ablation_value': re.compile(r'(消融|ablation).{0,120}(92\.2|88\.28)', re.I),
}


def extract_document(path):
    document = Document(path)
    paragraphs = [p.text.strip() for p in document.paragraphs if p.text.strip()]
    tables = []
    for table_index, table in enumerate(document.tables):
        rows = []
        for row in table.rows:
            rows.append([cell.text.strip() for cell in row.cells])
        tables.append({'index': table_index, 'rows': rows})
    return paragraphs, tables


def numbered_captions(paragraphs, kind):
    if kind == 'figure':
        pattern = re.compile(r'^(?:图|Figure|Fig\.?)\s*([0-9]+)', re.I)
    else:
        pattern = re.compile(r'^(?:表|Table)\s*([0-9]+)', re.I)
    found = []
    for index, text in enumerate(paragraphs):
        match = pattern.search(text)
        if match:
            found.append({'paragraph': index, 'number': int(match.group(1)), 'text': text})
    return found


def tpfs_rows(tables):
    rows = []
    for table in tables:
        for row_index, row in enumerate(table['rows']):
            joined = ' | '.join(row)
            if re.search(r'TPFS[-– ]?(Net|Former)', joined, re.I):
                rows.append({'table_index': table['index'], 'row_index': row_index, 'cells': row})
    return rows


def audit(path):
    paragraphs, tables = extract_document(path)
    text_units = paragraphs + [' | '.join(row) for table in tables for row in table['rows']]
    unsupported = []
    negated_seed_context = re.compile(
        r'(不是|不报告|尚未|待补|必须补|仅为?\s*1[- ]?seed|not\s+five)', re.I)
    for name, pattern in UNSUPPORTED_PATTERNS.items():
        for index, text in enumerate(text_units):
            if not pattern.search(text):
                continue
            if (name == 'unsupported_five_seed_robustness'
                    and negated_seed_context.search(text)):
                continue
            unsupported.append({'check': name, 'unit': index, 'text': text})
    figures = numbered_captions(paragraphs, 'figure')
    table_caps = numbered_captions(paragraphs, 'table')
    figure_numbers = [item['number'] for item in figures]
    table_numbers = [item['number'] for item in table_caps]
    return {
        'docx': str(Path(path).resolve()),
        'paragraph_count': len(paragraphs),
        'table_count': len(tables),
        'tpfs_rows': tpfs_rows(tables),
        'figure_captions': figures,
        'table_captions': table_caps,
        'figure_numbers_in_document_order': figure_numbers,
        'table_numbers_in_document_order': table_numbers,
        'figure_numbering_monotonic': figure_numbers == sorted(figure_numbers),
        'table_numbering_monotonic': table_numbers == sorted(table_numbers),
        'unsupported_claim_matches': unsupported,
    }


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('docx')
    parser.add_argument('--out')
    parser.add_argument('--strict', action='store_true')
    args = parser.parse_args()
    report = audit(args.docx)
    payload = json.dumps(report, ensure_ascii=False, indent=2)
    if args.out:
        out = Path(args.out)
        out.parent.mkdir(parents=True, exist_ok=True)
        out.write_text(payload)
        print(f'[saved] {out}')
    else:
        print(payload)
    failed = (
        not report['figure_numbering_monotonic']
        or not report['table_numbering_monotonic']
        or bool(report['unsupported_claim_matches'])
    )
    if args.strict and failed:
        raise SystemExit(1)


if __name__ == '__main__':
    main()
